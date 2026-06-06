#!/usr/bin/env python3
"""
Employee Scheduling Tool (Excel / openpyxl)
===========================================

Generates a single .xlsx workbook for a business open 7 days/week,
7:00 AM-11:00 PM (16 operating hours/day), scheduled in 30-minute blocks.

Sheets
------
    Inputs          Roster, operating hours, and all tunable parameters.
    Unavailability  Grid (employees x 30-min blocks); "X" = NOT available.
    Schedule        Auto-populated by the solver (static values, color-coded).
    Summary         Paid hours / days / shift lengths, per-block coverage,
                    feasibility check.
    README          How to set hours, mark unavailability, tune, re-run, and
                    how breaks affect paid time.

Break rules
-----------
    MEAL_BREAK_MIN  Unpaid, EXCLUDED from paid/work hours. Required for any
                    shift longer than 5.0 hrs. Placed near the middle of the
                    shift; the employee is on-premises but off the clock, so
                    that block provides NO coverage.
    REST_BREAK_MIN  Paid, COUNTED as work time. One per 4 hrs worked (or major
                    fraction). Marked distinctly but still covers its block.

Usage
-----
    python schedule_tool.py                 # build, or re-solve if file exists
    python schedule_tool.py --reset         # rebuild fresh sample data
    python schedule_tool.py --file plan.xlsx

All thresholds live in the Inputs sheet and are referenced by Excel defined
names, so Summary formulas stay editable. The assignment logic runs here in
Python and writes static values into Schedule.
"""

from __future__ import annotations

import argparse
import datetime as dt
import math
import os
import sys
from dataclasses import dataclass, field

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

DAYS = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]
WEEKEND = {5, 6}                 # Sat, Sun
EPS = 1e-9

# Cell marks used in the Schedule grid.
WORK = "●"      # paid working block
REST = "R"      # paid rest break (counts as work)
MEAL = "M"      # unpaid meal break (excluded from paid hours, no coverage)
OFF = "·"       # unavailable

MEAL_TRIGGER_HRS = 5.0          # shifts strictly longer than this need a meal

# ----------------------------------------------------------------------------
# Defaults (used only when seeding a brand-new workbook).
# ----------------------------------------------------------------------------
DEFAULT_EMPLOYEES = ["Alex", "Bina", "Cleo", "Dario"]
EMPLOYEE_ROWS = 8
DEFAULT_OPEN = dt.time(7, 0)     # 7:00 AM
DEFAULT_CLOSE = dt.time(23, 0)   # 11:00 PM
DEFAULT_PARAMS = {
    "WEEKDAY_STAFF": 2,
    "WEEKEND_STAFF": 2,
    "TARGET_MIN_WEEKLY": 30,
    "TARGET_MAX_WEEKLY": 35,
    "MAX_HOURS_PER_DAY": 8,
    "SHIFT_BLOCK_MINUTES": 30,
    "MEAL_BREAK_MIN": 30,
    "REST_BREAK_MIN": 10,
}
# Inputs label -> defined name used in formulas
PARAM_DEFNAME = {
    "WEEKDAY_STAFF": "WEEKDAY_STAFF",
    "WEEKEND_STAFF": "WEEKEND_STAFF",
    "TARGET_MIN_WEEKLY": "TARGET_MIN",
    "TARGET_MAX_WEEKLY": "TARGET_MAX",
    "MAX_HOURS_PER_DAY": "MAX_DAY",
    "SHIFT_BLOCK_MINUTES": "BLOCK_MIN",
    "MEAL_BREAK_MIN": "MEAL_MIN",
    "REST_BREAK_MIN": "REST_MIN",
}

EMP_COLORS = [
    "BBDEFB", "C8E6C9", "FFE0B2", "F8BBD0", "D1C4E9",
    "B2EBF2", "DCEDC8", "FFF59D", "FFCCBC", "E1BEE7",
]
UNAVAIL_FILL = PatternFill("solid", fgColor="78909C")
MEAL_FILL = PatternFill("solid", fgColor="EF6C00")       # orange = meal
REST_FILL = PatternFill("solid", fgColor="00897B")       # teal   = rest
HEADER_FILL = PatternFill("solid", fgColor="263238")
DAYHEAD_FILL = PatternFill("solid", fgColor="455A64")
FLAG_FILL = PatternFill("solid", fgColor="FFCDD2")
BAD_FILL = PatternFill("solid", fgColor="EF5350")
WHITE = Font(color="FFFFFF", bold=True)
BOLD = Font(bold=True)
CENTER = Alignment(horizontal="center", vertical="center")
WRAP = Alignment(horizontal="left", vertical="top", wrap_text=True)
THIN = Side(style="thin", color="B0BEC5")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


# ----------------------------------------------------------------------------
# Configuration model
# ----------------------------------------------------------------------------
@dataclass
class Block:
    day_idx: int
    slot: int
    start_min: int
    label: str


@dataclass
class Config:
    employees: list[str]
    open_min: list[int]
    close_min: list[int]
    params: dict[str, float]
    blocks: list[Block] = field(default_factory=list)
    day_slots: list[int] = field(default_factory=list)

    @property
    def block_minutes(self) -> int:
        return int(self.params["SHIFT_BLOCK_MINUTES"])

    @property
    def block_hours(self) -> float:
        return self.block_minutes / 60.0

    @property
    def meal_blocks(self) -> int:
        return max(1, round(self.params["MEAL_BREAK_MIN"] / self.block_minutes))

    @property
    def max_day_paid_blocks(self) -> int:
        return int(self.params["MAX_HOURS_PER_DAY"] / self.block_hours + EPS)

    @property
    def max_week_paid_blocks(self) -> int:
        return int(self.params["TARGET_MAX_WEEKLY"] / self.block_hours + EPS)

    @property
    def min_week_paid_blocks(self) -> int:
        return math.ceil(self.params["TARGET_MIN_WEEKLY"] / self.block_hours - EPS)

    def staff_needed(self, day_idx: int) -> int:
        key = "WEEKEND_STAFF" if day_idx in WEEKEND else "WEEKDAY_STAFF"
        return int(self.params[key])

    def needs_meal(self, span_blocks: int) -> bool:
        return span_blocks * self.block_hours > MEAL_TRIGGER_HRS + EPS

    def paid_blocks_of_span(self, span_blocks: int) -> int:
        if span_blocks <= 0:
            return 0
        return span_blocks - (self.meal_blocks if self.needs_meal(span_blocks) else 0)

    def build_blocks(self) -> None:
        self.blocks, self.day_slots = [], []
        step = self.block_minutes
        for d in range(7):
            slot, t = 0, self.open_min[d]
            while t + step <= self.close_min[d] + EPS:
                self.blocks.append(Block(d, slot, t, fmt_label(t)))
                slot += 1
                t += step
            self.day_slots.append(slot)

    def day_block_indices(self, day_idx: int) -> list[int]:
        return [i for i, b in enumerate(self.blocks) if b.day_idx == day_idx]


def fmt_label(minutes: int) -> str:
    h, m = divmod(minutes, 60)
    suffix = "a" if h < 12 else "p"
    hh = h % 12 or 12
    return f"{hh}{suffix}" if m == 0 else f"{hh}:{m:02d}{suffix}"


def fmt_time_long(minutes: int) -> str:
    h, m = divmod(minutes, 60)
    return f"{h % 12 or 12}:{m:02d} {'AM' if h < 12 else 'PM'}"


def minutes_to_time(minutes: int) -> dt.time:
    return dt.time(minutes // 60, minutes % 60)


# ----------------------------------------------------------------------------
# Solver: greedy/heuristic with one contiguous shift per day + breaks
# ----------------------------------------------------------------------------
@dataclass
class ShiftInfo:
    start: int          # first occupied slot (within day)
    span: int           # occupied blocks (on-premises), includes meal
    meal_slot: int      # day-slot of meal block, or -1
    rest_slots: list[int]   # day-slots marked as rest breaks
    paid_blocks: int


class Solver:
    def __init__(self, cfg: Config, unavail: dict[str, set[int]]):
        self.cfg = cfg
        self.unavail = unavail
        self.emps = cfg.employees
        self.run: dict[str, list] = {e: [None] * 7 for e in self.emps}  # (start,end)
        self.stuck: list[tuple[str, str]] = []
        self.shifts: dict[str, dict[int, ShiftInfo]] = {e: {} for e in self.emps}

    # geometry ---------------------------------------------------------------
    def _gid(self, day: int, slot: int) -> int:
        return self.cfg.day_block_indices(day)[slot]

    def is_available(self, name: str, day: int, slot: int) -> bool:
        return self._gid(day, slot) not in self.unavail.get(name, set())

    def span(self, name: str, day: int) -> int:
        a = self.run[name][day]
        return 0 if a is None else a[1] - a[0]

    def day_paid(self, name: str, day: int) -> int:
        return self.cfg.paid_blocks_of_span(self.span(name, day))

    def week_paid(self, name: str) -> int:
        return sum(self.day_paid(name, d) for d in range(7))

    def week_paid_hours(self, name: str) -> float:
        return self.week_paid(name) * self.cfg.block_hours

    # add logic --------------------------------------------------------------
    def can_add(self, name: str, day: int, slot: int) -> bool:
        if slot < 0 or slot >= self.cfg.day_slots[day]:
            return False
        if not self.is_available(name, day, slot):
            return False
        a = self.run[name][day]
        if a is None:
            new_span = 1
        elif slot == a[0] - 1 or slot == a[1]:
            new_span = (a[1] - a[0]) + 1
        else:
            return False
        new_paid = self.cfg.paid_blocks_of_span(new_span)
        if new_paid > self.cfg.max_day_paid_blocks:
            return False
        delta = new_paid - self.day_paid(name, day)
        if self.week_paid(name) + delta > self.cfg.max_week_paid_blocks:
            return False
        return True

    def do_add(self, name: str, day: int, slot: int) -> None:
        a = self.run[name][day]
        if a is None:
            self.run[name][day] = (slot, slot + 1)
        elif slot == a[0] - 1:
            self.run[name][day] = (slot, a[1])
        else:
            self.run[name][day] = (a[0], slot + 1)

    # paid-coverage of a (day, slot) given current runs ----------------------
    def covering(self, day: int, slot: int) -> int:
        c = 0
        for e in self.emps:
            a = self.run[e][day]
            if a is None or not (a[0] <= slot < a[1]):
                continue
            length = a[1] - a[0]
            if self.cfg.needs_meal(length):
                meal = a[0] + length // 2
                if slot == meal:
                    continue            # meal block: no coverage
            c += 1
        return c

    # feasibility bound: max paid blocks a person could reach ----------------
    def max_possible_blocks(self, name: str) -> int:
        total = 0
        for d in range(7):
            best = run = 0
            for slot in range(self.cfg.day_slots[d]):
                if self.is_available(name, d, slot):
                    run += 1
                    best = max(best, run)
                else:
                    run = 0
            # cap the usable span so paid blocks fit the daily cap
            span = best
            while self.cfg.paid_blocks_of_span(span) > self.cfg.max_day_paid_blocks:
                span -= 1
            total += self.cfg.paid_blocks_of_span(span)
        return min(total, self.cfg.max_week_paid_blocks)

    # main -------------------------------------------------------------------
    def solve(self) -> None:
        cfg = self.cfg

        # Phase A: staff each block up to its required count, balancing load.
        for blk in cfg.blocks:
            day, slot = blk.day_idx, blk.slot
            need = cfg.staff_needed(day)
            while self.covering(day, slot) < need:
                cands = [e for e in self.emps if self.can_add(e, day, slot)]
                if not cands:
                    break
                cands.sort(key=lambda e: (
                    self.week_paid(e) >= cfg.min_week_paid_blocks,
                    self.week_paid(e),
                    self.day_paid(e, day),
                ))
                self.do_add(cands[0], day, slot)

        # Phase B: bring everyone up to the weekly minimum, fairly.
        active = set(self.emps)
        while active:
            below = [e for e in active
                     if self.week_paid(e) < cfg.min_week_paid_blocks]
            if not below:
                break
            below.sort(key=self.week_paid)
            name = below[0]
            choice = self._best_fill_slot(name)
            if choice is None:
                active.discard(name)
                continue
            self.do_add(name, *choice)

        self._finalize_breaks()

        # Phase C: who fell short and why.
        for e in self.emps:
            if self.week_paid_hours(e) + EPS < cfg.params["TARGET_MIN_WEEKLY"]:
                cap = self.max_possible_blocks(e) * cfg.block_hours
                if cap + EPS < cfg.params["TARGET_MIN_WEEKLY"]:
                    reason = (f"availability caps paid time at {cap:g}h/week "
                              f"< {cfg.params['TARGET_MIN_WEEKLY']:g}h minimum")
                else:
                    reason = (f"reached {self.week_paid_hours(e):g}h; staffing "
                              f"demand was exhausted before the minimum "
                              f"(theoretical max {cap:g}h)")
                self.stuck.append((e, reason))

    def _best_fill_slot(self, name: str):
        best = best_key = None
        for d in range(7):
            need = self.cfg.staff_needed(d)
            for slot in range(self.cfg.day_slots[d]):
                if not self.can_add(name, d, slot):
                    continue
                short = max(0, need - self.covering(d, slot))   # coverage gap
                has = 0 if self.run[name][d] is None else 1
                key = (-short, -has, self._gid(d, slot))
                if best_key is None or key < best_key:
                    best_key, best = key, (d, slot)
        return best

    def _finalize_breaks(self) -> None:
        cfg = self.cfg
        for e in self.emps:
            self.shifts[e] = {}
            for d in range(7):
                a = self.run[e][d]
                if a is None:
                    continue
                start, end = a
                length = end - start
                meal_slot = -1
                if cfg.needs_meal(length):
                    meal_slot = start + length // 2
                paid = cfg.paid_blocks_of_span(length)
                paid_hours = paid * cfg.block_hours
                # rest breaks: one per 4 paid hrs (or major fraction >= 2h)
                whole = int(paid_hours // 4)
                rest_n = whole + (1 if (paid_hours - 4 * whole) >= 2 - EPS else 0)
                # place rest marks on paid blocks, evenly, never on the meal
                paid_slots = [s for s in range(start, end) if s != meal_slot]
                rest_slots = []
                if rest_n and paid_slots:
                    for k in range(1, rest_n + 1):
                        idx = min(len(paid_slots) - 1,
                                  int(k * len(paid_slots) / (rest_n + 1)))
                        s = paid_slots[idx]
                        if s not in rest_slots:
                            rest_slots.append(s)
                self.shifts[e][d] = ShiftInfo(start, length, meal_slot,
                                              rest_slots, paid)

    # export helpers ---------------------------------------------------------
    def cell_kind(self, name: str, day: int, slot: int) -> str:
        if not self.is_available(name, day, slot):
            return OFF
        info = self.shifts[name].get(day)
        if not info or not (info.start <= slot < info.start + info.span):
            return ""
        if slot == info.meal_slot:
            return MEAL
        if slot in info.rest_slots:
            return REST
        return WORK


# ----------------------------------------------------------------------------
# Validation
# ----------------------------------------------------------------------------
def validate(solver: Solver) -> list[str]:
    cfg = solver.cfg
    errs: list[str] = []
    for e in solver.emps:
        wk = solver.week_paid_hours(e)
        if wk - EPS > cfg.params["TARGET_MAX_WEEKLY"]:
            errs.append(f"{e}: weekly paid {wk:g}h exceeds TARGET_MAX_WEEKLY")
        for d in range(7):
            info = solver.shifts[e].get(d)
            if not info:
                continue
            # contiguity is guaranteed by construction; verify availability
            for slot in range(info.start, info.start + info.span):
                if not solver.is_available(e, d, slot):
                    errs.append(f"{e}: scheduled while unavailable on {DAYS[d]}")
            # daily paid cap
            if info.paid_blocks * cfg.block_hours - EPS > cfg.params["MAX_HOURS_PER_DAY"]:
                errs.append(f"{e}: {DAYS[d]} paid hours exceed MAX_HOURS_PER_DAY")
            # meal required & excluded for shifts > 5h
            if cfg.needs_meal(info.span):
                if info.meal_slot < 0:
                    errs.append(f"{e}: {DAYS[d]} shift >5h has no meal break")
                if info.paid_blocks != info.span - cfg.meal_blocks:
                    errs.append(f"{e}: {DAYS[d]} meal not excluded from paid hours")
            elif info.meal_slot >= 0:
                errs.append(f"{e}: {DAYS[d]} meal break on a shift <=5h")
            # rest breaks counted as work => paid blocks include them
            if any(s == info.meal_slot for s in info.rest_slots):
                errs.append(f"{e}: {DAYS[d]} rest break overlaps meal break")
    return errs


# ----------------------------------------------------------------------------
# Feasibility check
# ----------------------------------------------------------------------------
def feasibility(cfg: Config) -> dict:
    op_hours = sum((cfg.close_min[d] - cfg.open_min[d]) / 60.0 for d in range(7))
    required = sum((cfg.close_min[d] - cfg.open_min[d]) / 60.0
                   * cfg.staff_needed(d) for d in range(7))
    per_emp_cap = min(cfg.params["MAX_HOURS_PER_DAY"] * 7,
                      cfg.params["TARGET_MAX_WEEKLY"])
    n = len(cfg.employees)
    supply = n * per_emp_cap
    feasible = supply + EPS >= required
    emps_needed = math.ceil(required / per_emp_cap) if per_emp_cap else 0
    # weekly hours each of the current n would need (caps aside)
    weekly_each_needed = required / n if n else float("inf")
    daily_each_needed = required / (n * 7) if n else float("inf")
    return {
        "operating_hours_week": op_hours,
        "required_coverage_hours": required,
        "per_emp_cap": per_emp_cap,
        "supply": supply,
        "feasible": feasible,
        "emps_needed": emps_needed,
        "weekly_each_needed": weekly_each_needed,
        "daily_each_needed": daily_each_needed,
        "binding_weekly": cfg.params["TARGET_MAX_WEEKLY"] < cfg.params["MAX_HOURS_PER_DAY"] * 7,
    }


def feasibility_lines(cfg: Config, fb: dict, solver: Solver) -> list[str]:
    p = cfg.params
    lines = [
        f"Operating coverage required: {fb['required_coverage_hours']:g} "
        f"staff-hours/week "
        f"({sum((cfg.close_min[d]-cfg.open_min[d])/60 for d in range(7))/7:g}h/day "
        f"x staff needed x 7 days).",
        f"Per-employee supply cap: {fb['per_emp_cap']:g} h/week "
        f"(min of MAX_HOURS_PER_DAY x 7 = {p['MAX_HOURS_PER_DAY']*7:g} and "
        f"TARGET_MAX_WEEKLY = {p['TARGET_MAX_WEEKLY']:g}).",
        f"With {len(cfg.employees)} employees, total supply = "
        f"{fb['supply']:g} h/week.",
    ]
    if fb["feasible"]:
        lines.append("=> FEASIBLE: total supply meets the required coverage "
                     "(individual availability gaps may still leave blocks "
                     "uncovered).")
    else:
        lines.append(
            f"=> INFEASIBLE: supply {fb['supply']:g}h < required "
            f"{fb['required_coverage_hours']:g}h. Full 2-person coverage of "
            f"{fb['operating_hours_week']/7:g} operating hours cannot be met.")
        lines.append(
            f"   To make it feasible: use at least {fb['emps_needed']} "
            f"employees at the current caps, OR raise TARGET_MAX_WEEKLY to "
            f">= {math.ceil(fb['weekly_each_needed'])}h so the existing "
            f"{len(cfg.employees)} can each cover "
            f"{fb['weekly_each_needed']:g}h/week "
            f"(~{fb['daily_each_needed']:g}h/day, 7 days).")
        if fb["binding_weekly"]:
            lines.append(
                "   Note: TARGET_MAX_WEEKLY is the binding constraint here, so "
                "raising MAX_HOURS_PER_DAY alone will NOT help.")
    return lines


# ----------------------------------------------------------------------------
# Sample data
# ----------------------------------------------------------------------------
def seed_unavailability(cfg: Config) -> dict[str, set[int]]:
    u: dict[str, set[int]] = {e: set() for e in cfg.employees}

    def mark(name, day_idx, predicate):
        for gid in cfg.day_block_indices(day_idx):
            if predicate(cfg.blocks[gid].start_min):
                u[name].add(gid)

    for d in range(7):
        # Alex: no early mornings before 9a.
        mark("Alex", d, lambda m: m < 9 * 60)
        # Cleo: no late evenings from 8p onward.
        mark("Cleo", d, lambda m: m >= 20 * 60)
    # Bina: weekends fully off (likely short of the weekly minimum).
    for d in (5, 6):
        for gid in cfg.day_block_indices(d):
            u["Bina"].add(gid)
    # Dario: Wednesdays off; no blocks before 11a on Mondays.
    for gid in cfg.day_block_indices(2):
        u["Dario"].add(gid)
    mark("Dario", 0, lambda m: m < 11 * 60)
    return u


def default_config() -> Config:
    cfg = Config(
        employees=list(DEFAULT_EMPLOYEES),
        open_min=[DEFAULT_OPEN.hour * 60 + DEFAULT_OPEN.minute] * 7,
        close_min=[DEFAULT_CLOSE.hour * 60 + DEFAULT_CLOSE.minute] * 7,
        params=dict(DEFAULT_PARAMS),
    )
    cfg.build_blocks()
    return cfg


# ----------------------------------------------------------------------------
# Reading an existing workbook back in
# ----------------------------------------------------------------------------
def _cell_to_minutes(value, default: int) -> int:
    if value is None:
        return default
    if isinstance(value, (dt.time, dt.datetime)):
        return value.hour * 60 + value.minute
    if isinstance(value, (int, float)):
        v = float(value)
        return int(round(v * 1440)) if 0 < v < 1 else int(round(v * 60))
    s = str(value).strip()
    for fmt in ("%I:%M %p", "%H:%M", "%I %p"):
        try:
            t = dt.datetime.strptime(s, fmt).time()
            return t.hour * 60 + t.minute
        except ValueError:
            continue
    return default


def read_config(wb) -> Config:
    ws = wb["Inputs"]
    labels = {}
    for row in ws.iter_rows(min_col=1, max_col=1):
        c = row[0]
        if isinstance(c.value, str):
            labels[c.value.strip()] = c.row

    employees = []
    if "Employees" in labels:
        rr = labels["Employees"] + 2
        while ws.cell(rr, 1).value not in (None, ""):
            name = str(ws.cell(rr, 1).value).strip()
            if name:
                employees.append(name)
            rr += 1
    if not employees:
        employees = list(DEFAULT_EMPLOYEES)

    open_min = [DEFAULT_OPEN.hour * 60] * 7
    close_min = [DEFAULT_CLOSE.hour * 60] * 7
    if "Operating hours" in labels:
        r = labels["Operating hours"]
        for d in range(7):
            row = r + 2 + d
            open_min[d] = _cell_to_minutes(ws.cell(row, 2).value, open_min[d])
            close_min[d] = _cell_to_minutes(ws.cell(row, 3).value, close_min[d])

    params = dict(DEFAULT_PARAMS)
    for key in params:
        if key in labels:
            v = ws.cell(labels[key], 2).value
            if isinstance(v, (int, float)):
                params[key] = v

    cfg = Config(employees, open_min, close_min, params)
    cfg.build_blocks()
    return cfg


def read_unavailability(wb, cfg: Config) -> dict[str, set[int]]:
    u = {e: set() for e in cfg.employees}
    if "Unavailability" not in wb.sheetnames:
        return u
    ws = wb["Unavailability"]
    for row in ws.iter_rows(min_row=3, min_col=1):
        name = row[0].value
        if not name:
            continue
        name = str(name).strip()
        if name not in u:
            continue
        for gid in range(len(cfg.blocks)):
            if ws.cell(row[0].row, 2 + gid).value not in (None, ""):
                u[name].add(gid)
    return u


# ----------------------------------------------------------------------------
# Workbook writers
# ----------------------------------------------------------------------------
def style_header(cell, text, fill=HEADER_FILL):
    cell.value = text
    cell.fill = fill
    cell.font = WHITE
    cell.alignment = CENTER
    cell.border = BORDER


def write_inputs(wb, cfg: Config) -> None:
    ws = wb.create_sheet("Inputs")
    ws.sheet_properties.tabColor = "1565C0"
    ws.column_dimensions["A"].width = 26
    for col in "BCD":
        ws.column_dimensions[col].width = 14

    ws["A1"] = "Employee Scheduling - Inputs"
    ws["A1"].font = Font(bold=True, size=14)

    r = 3
    ws.cell(r, 1, "Employees").font = BOLD
    style_header(ws.cell(r + 1, 1), "Name")
    names = list(cfg.employees) + [""] * max(0, EMPLOYEE_ROWS - len(cfg.employees))
    for i, name in enumerate(names):
        ws.cell(r + 2 + i, 1, name).border = BORDER
    bh = r + 2 + len(names) + 1

    ws.cell(bh, 1, "Operating hours").font = BOLD
    style_header(ws.cell(bh + 1, 1), "Day")
    style_header(ws.cell(bh + 1, 2), "Open")
    style_header(ws.cell(bh + 1, 3), "Close")
    for d in range(7):
        row = bh + 2 + d
        ws.cell(row, 1, DAYS[d]).border = BORDER
        for col, mins in ((2, cfg.open_min[d]), (3, cfg.close_min[d])):
            c = ws.cell(row, col, minutes_to_time(mins))
            c.number_format = "h:mm AM/PM"
            c.border = BORDER
            c.alignment = CENTER

    pr = bh + 2 + 7 + 1
    ws.cell(pr, 1, "Parameters").font = BOLD
    style_header(ws.cell(pr + 1, 1), "Parameter")
    style_header(ws.cell(pr + 1, 2), "Value")
    refs = {}
    for i, (key, defname) in enumerate(PARAM_DEFNAME.items()):
        row = pr + 2 + i
        ws.cell(row, 1, key).border = BORDER
        vc = ws.cell(row, 2, cfg.params[key])
        vc.border = BORDER
        vc.alignment = CENTER
        vc.font = BOLD
        refs[defname] = f"Inputs!$B${row}"

    note = pr + 2 + len(PARAM_DEFNAME) + 1
    ws.cell(note, 1,
            "Edit any value above, then re-run schedule_tool.py to refresh the "
            "Schedule and Summary sheets. Meal breaks are unpaid (excluded from "
            "paid hours); rest breaks are paid (counted as work).").font = \
        Font(italic=True, size=9)

    for defname, ref in refs.items():
        if defname in wb.defined_names:
            del wb.defined_names[defname]
        wb.defined_names.add(DefinedName(defname, attr_text=ref))


def write_grid_header(ws, cfg: Config, corner: str) -> int:
    style_header(ws.cell(1, 1), corner, DAYHEAD_FILL)
    style_header(ws.cell(2, 1), "Employee \\ Block", DAYHEAD_FILL)
    ws.column_dimensions["A"].width = 16
    col = 2
    for d in range(7):
        idxs = cfg.day_block_indices(d)
        if not idxs:
            continue
        start_col = col
        for gid in idxs:
            style_header(ws.cell(2, col), cfg.blocks[gid].label, HEADER_FILL)
            ws.column_dimensions[get_column_letter(col)].width = 4.5
            col += 1
        ws.merge_cells(start_row=1, start_column=start_col,
                       end_row=1, end_column=col - 1)
        style_header(ws.cell(1, start_col), DAYS[d], DAYHEAD_FILL)
    return col - 1


def write_unavailability(wb, cfg: Config, unavail: dict[str, set[int]]):
    ws = wb.create_sheet("Unavailability")
    ws.sheet_properties.tabColor = "C62828"
    last_col = write_grid_header(ws, cfg, "Unavailability")
    ws.freeze_panes = "B3"

    for i, name in enumerate(cfg.employees):
        row = 3 + i
        nc = ws.cell(row, 1, name)
        nc.font = BOLD
        nc.border = BORDER
        for gid in range(len(cfg.blocks)):
            cell = ws.cell(row, 2 + gid)
            cell.border = BORDER
            cell.alignment = CENTER
            if gid in unavail.get(name, set()):
                cell.value = "X"
    last_row = 2 + len(cfg.employees)

    rng = f"B3:{get_column_letter(last_col)}{last_row}"
    dv = DataValidation(type="list", formula1='"X,x"', allow_blank=True,
                        showErrorMessage=True)
    dv.error = 'Enter "X" to mark unavailable, or leave blank for available.'
    dv.errorTitle = "Invalid entry"
    ws.add_data_validation(dv)
    dv.add(rng)
    ws.conditional_formatting.add(
        rng, FormulaRule(formula=['UPPER(B3)="X"'],
                         fill=PatternFill("solid", fgColor="EF9A9A")))
    ws.cell(last_row + 2, 1,
            'Mark "X" where an employee is NOT available; blank = available. '
            "Re-run schedule_tool.py after editing.").font = \
        Font(italic=True, size=9)


def write_schedule(wb, cfg: Config, solver: Solver):
    ws = wb.create_sheet("Schedule")
    ws.sheet_properties.tabColor = "2E7D32"
    last_col = write_grid_header(ws, cfg, "Schedule")
    ws.freeze_panes = "B3"

    color_of = {e: EMP_COLORS[i % len(EMP_COLORS)]
                for i, e in enumerate(cfg.employees)}

    for i, name in enumerate(cfg.employees):
        row = 3 + i
        nc = ws.cell(row, 1, name)
        nc.font = BOLD
        nc.border = BORDER
        nc.fill = PatternFill("solid", fgColor=color_of[name])
        for d in range(7):
            for slot, gid in enumerate(cfg.day_block_indices(d)):
                cell = ws.cell(row, 2 + gid)
                cell.border = BORDER
                cell.alignment = CENTER
                kind = solver.cell_kind(name, d, slot)
                if kind == OFF:
                    cell.value = OFF
                    cell.fill = UNAVAIL_FILL
                    cell.font = Font(color="ECEFF1")
                elif kind == WORK:
                    cell.value = WORK
                    cell.fill = PatternFill("solid", fgColor=color_of[name])
                    cell.font = Font(color=color_of[name])
                elif kind == REST:
                    cell.value = REST
                    cell.fill = REST_FILL
                    cell.font = WHITE
                elif kind == MEAL:
                    cell.value = MEAL
                    cell.fill = MEAL_FILL
                    cell.font = WHITE
    last_row = 2 + len(cfg.employees)

    # Coverage row (paid presence per block) + required comparison.
    cov_row = last_row + 1
    req_row = last_row + 2
    for label, rr in (("Scheduled", cov_row), ("Required", req_row)):
        c = ws.cell(rr, 1, label)
        c.fill = DAYHEAD_FILL
        c.font = WHITE
        c.border = BORDER
    for gid, blk in enumerate(cfg.blocks):
        cov = solver.covering(blk.day_idx, blk.slot)
        need = cfg.staff_needed(blk.day_idx)
        cc = ws.cell(cov_row, 2 + gid, cov)
        rc = ws.cell(req_row, 2 + gid, need)
        for cell in (cc, rc):
            cell.alignment = CENTER
            cell.border = BORDER
            cell.font = BOLD
    ws.conditional_formatting.add(
        f"B{cov_row}:{get_column_letter(last_col)}{cov_row}",
        FormulaRule(formula=[f"B{cov_row}<B{req_row}"], fill=BAD_FILL))

    # Legend
    leg = req_row + 2
    ws.cell(leg, 1, "Legend").font = BOLD
    rowi = leg + 1
    for name in cfg.employees:
        c = ws.cell(rowi, 1, f"{name} (work {WORK})")
        c.fill = PatternFill("solid", fgColor=color_of[name])
        c.border = BORDER
        rowi += 1
    for text, fill, font in (("Rest break (R) - paid", REST_FILL, WHITE),
                             ("Meal break (M) - unpaid", MEAL_FILL, WHITE),
                             ("Unavailable", UNAVAIL_FILL,
                              Font(color="FFFFFF"))):
        c = ws.cell(rowi, 1, text)
        c.fill = fill
        c.font = font
        c.border = BORDER
        rowi += 1


def write_summary(wb, cfg: Config, solver: Solver, fb: dict):
    ws = wb.create_sheet("Summary")
    ws.sheet_properties.tabColor = "6A1B9A"
    ws.column_dimensions["A"].width = 18
    for col in "BCDE":
        ws.column_dimensions[col].width = 14
    ws.column_dimensions["F"].width = 40

    ws["A1"] = "Schedule Summary"
    ws["A1"].font = Font(bold=True, size=14)

    n = len(cfg.employees)
    nb = len(cfg.blocks)
    first = get_column_letter(2)
    last = get_column_letter(1 + nb)

    # Per-employee table -----------------------------------------------------
    r = 3
    ws.cell(r, 1, "Per-employee totals (paid hours)").font = BOLD
    for j, h in enumerate(["Employee", "Paid hours", "Days worked",
                           "Status", "Shift lengths (paid h)"]):
        style_header(ws.cell(r + 1, 1 + j), h)
    for i, name in enumerate(cfg.employees):
        row = r + 2 + i
        srow = 3 + i
        ws.cell(row, 1, name).border = BORDER
        rng = f"Schedule!{first}{srow}:{last}{srow}"
        # Paid hours = (work + rest) blocks * block hours; meal excluded.
        hcell = ws.cell(row, 2)
        hcell.value = (f'=(COUNTIF({rng},"{WORK}")+COUNTIF({rng},"{REST}"))'
                       f'*(BLOCK_MIN/60)')
        hcell.border = BORDER
        hcell.alignment = CENTER
        # Days worked + max daily paid, per day.
        day_terms, max_terms = [], []
        for d in range(7):
            idxs = cfg.day_block_indices(d)
            if not idxs:
                continue
            c0 = get_column_letter(2 + idxs[0])
            c1 = get_column_letter(2 + idxs[-1])
            dr = f"Schedule!{c0}{srow}:{c1}{srow}"
            cnt = f'(COUNTIF({dr},"{WORK}")+COUNTIF({dr},"{REST}"))'
            day_terms.append(f"({cnt}>0)")
            max_terms.append(cnt)
        ws.cell(row, 3, "=" + "+".join(day_terms)).border = BORDER
        ws.cell(row, 3).alignment = CENTER
        maxday = f'(MAX({",".join(max_terms)})*(BLOCK_MIN/60))'
        hc = f"B{row}"
        scell = ws.cell(row, 4)
        scell.value = (f'=IF({hc}<TARGET_MIN,"BELOW MIN",'
                       f'IF({hc}>TARGET_MAX,"ABOVE MAX",'
                       f'IF({maxday}>MAX_DAY,"OVER DAILY CAP","OK")))')
        scell.border = BORDER
        scell.alignment = CENTER
        # Shift lengths (static, descriptive).
        parts = []
        for d in range(7):
            info = solver.shifts[name].get(d)
            if info:
                parts.append(f"{DAYS[d]} {info.paid_blocks*cfg.block_hours:g}")
        lc = ws.cell(row, 5, ", ".join(parts) if parts else "-")
        lc.border = BORDER
        lc.alignment = Alignment(horizontal="left", vertical="center")
    emp_last = r + 1 + n
    ws.conditional_formatting.add(
        f"D{r+2}:D{emp_last}",
        FormulaRule(formula=[f'D{r+2}<>"OK"'], fill=FLAG_FILL))

    # Coverage table ---------------------------------------------------------
    cov_start = emp_last + 2
    ws.cell(cov_start, 1, "Coverage per time block").font = BOLD
    for j, h in enumerate(["Day", "Time", "Scheduled", "Required", "Status"]):
        style_header(ws.cell(cov_start + 1, 1 + j), h)
    for gid, blk in enumerate(cfg.blocks):
        row = cov_start + 2 + gid
        ws.cell(row, 1, DAYS[blk.day_idx]).border = BORDER
        ws.cell(row, 2, fmt_time_long(blk.start_min)).border = BORDER
        col = get_column_letter(2 + gid)
        # paid presence excludes meal (M) blocks automatically.
        sc = ws.cell(row, 3)
        sc.value = (f'=COUNTIF(Schedule!{col}3:{col}{2+n},"{WORK}")'
                    f'+COUNTIF(Schedule!{col}3:{col}{2+n},"{REST}")')
        sc.border = BORDER
        sc.alignment = CENTER
        need_name = "WEEKEND_STAFF" if blk.day_idx in WEEKEND else "WEEKDAY_STAFF"
        rc = ws.cell(row, 4, f"={need_name}")
        rc.border = BORDER
        rc.alignment = CENTER
        st = ws.cell(row, 5, f'=IF(C{row}<D{row},"UNCOVERED","ok")')
        st.border = BORDER
        st.alignment = CENTER
    cov_last = cov_start + 1 + nb
    ws.conditional_formatting.add(
        f"C{cov_start+2}:E{cov_last}",
        FormulaRule(formula=[f"$C{cov_start+2}<$D{cov_start+2}"], fill=FLAG_FILL))

    # Feasibility + notes ----------------------------------------------------
    fr = cov_last + 2
    ws.cell(fr, 1, "Feasibility check").font = BOLD
    lines = feasibility_lines(cfg, fb, solver)
    lines.append("")
    if solver.stuck:
        lines.append("Employees who cannot reach TARGET_MIN_WEEKLY within "
                     "their availability:")
        for name, reason in solver.stuck:
            lines.append(f"  - {name}: {reason}")
    else:
        lines.append("All employees reached the weekly minimum.")
    lines.append("")
    lines.append("Breaks: meal blocks (M) are unpaid and excluded from paid "
                 "hours and from coverage (an uncovered gap can appear there if "
                 "no co-worker overlaps). Rest blocks (R) are paid, count as "
                 "work, and do cover their block.")
    for i, text in enumerate(lines):
        c = ws.cell(fr + 1 + i, 1, text)
        c.alignment = WRAP
        ws.merge_cells(start_row=fr + 1 + i, start_column=1,
                       end_row=fr + 1 + i, end_column=6)


def write_readme(wb):
    ws = wb.create_sheet("README")
    ws.sheet_properties.tabColor = "F9A825"
    ws.column_dimensions["A"].width = 100
    ws["A1"] = "Employee Scheduling Tool - README"
    ws["A1"].font = Font(bold=True, size=14)
    steps = [
        "",
        "OVERVIEW",
        "  Plans weekly shifts for a business open 7 days/week in 30-minute "
        "blocks. You provide employees, operating hours, parameters, and each "
        "person's unavailability; the solver fills Schedule and writes Summary.",
        "",
        "SHEETS",
        "  - Inputs         roster, operating hours, tunable parameters.",
        "  - Unavailability employees x blocks; mark \"X\" where unavailable.",
        "  - Schedule       auto-filled, color-coded per employee.",
        "  - Summary        paid hours/days/shift lengths, coverage, feasibility.",
        "",
        "SET OPERATING HOURS",
        "  On Inputs, edit Open/Close per day. Blocks regenerate from these "
        "times and SHIFT_BLOCK_MINUTES.",
        "",
        "MARK UNAVAILABILITY",
        "  On Unavailability, type \"X\" in any block a person cannot work; "
        "blank means available. Cells accept only blank or X; X is shaded red.",
        "",
        "TUNE PARAMETERS (Inputs)",
        "  WEEKDAY_STAFF / WEEKEND_STAFF  staff required per block.",
        "  TARGET_MIN_WEEKLY / TARGET_MAX_WEEKLY  paid-hour band per person.",
        "  MAX_HOURS_PER_DAY  daily paid-hour cap.",
        "  SHIFT_BLOCK_MINUTES  block granularity (default 30).",
        "  MEAL_BREAK_MIN  unpaid meal length (excluded from paid hours).",
        "  REST_BREAK_MIN  paid rest length (counted as work).",
        "  Summary flags reference these cells, so nothing is hardcoded.",
        "",
        "HOW BREAKS AFFECT PAID TIME",
        "  Meal (M): shifts longer than 5.0 hrs get one unpaid meal near the "
        "middle. It is OFF the clock - excluded from paid hours and it provides "
        "NO coverage (a gap may appear unless a co-worker overlaps).",
        "  Rest (R): one paid 10-min rest per 4 hrs worked (or major fraction). "
        "Rest counts as work time and covers its block.",
        "  Paid hours = (work blocks + rest blocks) x block hours; meal excluded.",
        "",
        "RE-RUN",
        "  Save edits in Excel, then run:",
        "        python schedule_tool.py            (re-solve existing file)",
        "        python schedule_tool.py --reset    (rebuild sample data)",
        "  The script reads Inputs + Unavailability and overwrites Schedule + "
        "Summary with fresh static values, and prints a feasibility report and "
        "anyone who cannot be fully scheduled.",
        "",
        "CONSTRAINTS ENFORCED",
        "  1. Never scheduled in an \"X\" block.",
        "  2. Daily paid hours <= MAX_HOURS_PER_DAY.",
        "  3. Weekly paid hours within TARGET_MIN..TARGET_MAX when feasible.",
        "  4. One contiguous shift per day (no split shifts).",
        "  5. Each block staffed to the required count where availability allows.",
        "  6. Meal excluded from paid hours; rest included.",
    ]
    for i, text in enumerate(steps):
        c = ws.cell(2 + i, 1, text)
        c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        if text and text == text.upper() and not text.startswith(" "):
            c.font = BOLD


# ----------------------------------------------------------------------------
# Orchestration
# ----------------------------------------------------------------------------
def build_workbook(cfg, unavail, solver, fb, path):
    wb = Workbook()
    wb.remove(wb.active)
    write_inputs(wb, cfg)
    write_unavailability(wb, cfg, unavail)
    write_schedule(wb, cfg, solver)
    write_summary(wb, cfg, solver, fb)
    write_readme(wb)
    order = ["README", "Inputs", "Unavailability", "Schedule", "Summary"]
    wb._sheets.sort(key=lambda s: order.index(s.title))
    wb.active = wb.sheetnames.index("Schedule")
    wb.save(path)


def run(path: str, reset: bool):
    fresh = reset or not os.path.exists(path)
    if fresh:
        print(f"Building new workbook with sample data -> {path}")
        cfg = default_config()
        unavail = seed_unavailability(cfg)
    else:
        print(f"Reading existing workbook -> {path}")
        wb = load_workbook(path)
        cfg = read_config(wb)
        unavail = read_unavailability(wb, cfg)

    fb = feasibility(cfg)
    solver = Solver(cfg, unavail)
    solver.solve()
    errs = validate(solver)

    print("\n=== Feasibility check ===")
    for line in feasibility_lines(cfg, fb, solver):
        print(" ", line)

    print("\n=== Constraint validation ===")
    if errs:
        print("VIOLATIONS DETECTED:")
        for e in errs:
            print("  !", e)
    else:
        print("OK - no hard constraint violated (no work during X, daily cap, "
              "weekly max, contiguity, meal excluded, rest included).")

    print("\n=== Weekly paid hours per employee ===")
    for e in cfg.employees:
        days = sum(1 for d in range(7) if d in solver.shifts[e])
        print(f"  {e:<8} {solver.week_paid_hours(e):5.1f} paid h   {days} day(s)")

    print("\n=== Employees short of TARGET_MIN_WEEKLY ===")
    if solver.stuck:
        for name, reason in solver.stuck:
            print(f"  - {name}: {reason}")
    else:
        print("  none.")

    uncovered = [b for b in cfg.blocks
                 if solver.covering(b.day_idx, b.slot) < cfg.staff_needed(b.day_idx)]
    print(f"\n=== Coverage === blocks={len(cfg.blocks)} "
          f"uncovered={len(uncovered)}")
    if uncovered:
        sample = ", ".join(f"{DAYS[b.day_idx]} {b.label}" for b in uncovered[:10])
        more = "" if len(uncovered) <= 10 else f" (+{len(uncovered)-10} more)"
        print(f"  first uncovered: {sample}{more}")

    build_workbook(cfg, unavail, solver, fb, path)
    print(f"\nSaved {path}")
    return errs


def main():
    ap = argparse.ArgumentParser(description="Employee scheduling tool (Excel).")
    ap.add_argument("--file", default="employee_schedule.xlsx")
    ap.add_argument("--reset", action="store_true")
    args = ap.parse_args()
    errs = run(args.file, args.reset)
    sys.exit(1 if errs else 0)


if __name__ == "__main__":
    main()
